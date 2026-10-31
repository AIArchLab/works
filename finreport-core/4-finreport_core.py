#!/usr/bin/env python3
"""FinReport-Core MVP v0.2 - merge heterogeneous-format reports with field mapping,
normalization, and fail-closed validation. Pure rule engine; no network calls, no LLM API.
Third-party dependency: openpyxl only (for .xlsx I/O; the CSV path uses the standard library).
Usage: python3 finreport_core.py <mapping.json> <fileA> <fileB> <out_summary.xlsx> <error_list.csv>
Exit codes: 0=success, summary written; 1=usage error or output destination already exists
(overwrite refused); 2=validation blocked (error list written, no summary).
Safety: never overwrites or deletes existing files - an existing destination is refused,
so stale results can never be mistaken for this run's output."""
import csv, json, math, os, sys, datetime

DATE_FORMATS = ["%Y/%m/%d", "%Y-%m-%d", "%Y.%m.%d"]

def parse_date(v):
    s = str(v).strip()
    for fmt in DATE_FORMATS:
        try:
            datetime.datetime.strptime(s, fmt)
            return True
        except ValueError:
            pass
    return False

def parse_amount(v):
    try:
        x = float(str(v).replace(",", "").strip())
    except (ValueError, TypeError):
        return None
    return x if math.isfinite(x) else None  # reject nan/inf and other non-finite amounts

def read_rows(path):
    if path.lower().endswith(".xlsx"):
        from openpyxl import load_workbook
        ws = load_workbook(path, read_only=True).active
        rows = list(ws.iter_rows(values_only=True))
        header = [str(c).strip() for c in rows[0]]
        return header, [dict(zip(header, r)) for r in rows[1:]]
    with open(path, newline="", encoding="utf-8-sig") as f:
        rdr = csv.DictReader(f)
        return rdr.fieldnames, list(rdr)

def validate(rows, mapping, role, errors, reject_dup_keys):
    """Validate key/date/amount per mapping; reject_dup_keys=True flags duplicate keys."""
    out, seen = [], set()
    for i, row in enumerate(rows, start=2):  # row 1 is the header
        rec = {canon: row.get(src) for canon, src in mapping.items()}
        key = str(rec.get("agent_id") or "").strip()
        if not key:
            errors.append((role, i, "missing key (agent_id)", json.dumps(row, ensure_ascii=False)))
            continue
        if reject_dup_keys and key in seen:
            errors.append((role, i, "duplicate key (agent_id)", key))
            continue
        seen.add(key)
        if not parse_date(rec.get("date")):
            errors.append((role, i, "bad date format", str(rec.get("date"))))
            continue
        amt = parse_amount(rec.get("amount"))
        if amt is None:
            errors.append((role, i, "bad amount format (incl. non-finite)", str(rec.get("amount"))))
            continue
        out.append({"agent_id": key, "date": str(rec["date"]).strip(), "amount": amt})
    return out

def main():
    if len(sys.argv) != 6:
        print(__doc__); return 1
    _, mp, fa, fb, out_xlsx, err_csv = sys.argv
    for dest in (out_xlsx, err_csv):
        if os.path.exists(dest):
            print(f"refused: output destination already exists ({dest}). "
              "To avoid stale results being mistaken for this run's output, specify a new "
              "path; this tool never overwrites or deletes existing files.")
            return 1
    mapping = json.load(open(mp, encoding="utf-8"))
    errors, data = [], {}
    for role, path in (("A", fa), ("B", fb)):
        _, rows = read_rows(path)
        data[role] = validate(rows, mapping[role], role, errors, reject_dup_keys=(role == "A"))
    if errors:  # fail-closed: no summary when input is bad
        with open(err_csv, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow(["file", "row", "reason", "raw content"])
            w.writerows(errors)
        print(f"blocked: {len(errors)} issue(s) written to {err_csv}; no summary produced.")
        return 2
    agg = {}
    for r in data["B"]:
        agg[r["agent_id"]] = agg.get(r["agent_id"], 0.0) + r["amount"]  # detail rows aggregate by key, by design
    a_map = {r["agent_id"]: r["amount"] for r in data["A"]}
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active
    ws.append(["Agent ID", "Summary amount (A)", "Detail total (B)", "Diff (B-A)", "Status"])
    for aid in sorted(set(a_map) | set(agg)):
        a, b = a_map.get(aid), agg.get(aid)
        if a is None: st, diff = "detail only", b
        elif b is None: st, diff = "summary only", -a
        else:
            diff = round(b - a, 2)
            st = "OK" if abs(diff) < 0.005 else "MISMATCH"
        ws.append([aid, a if a is not None else "", b if b is not None else "", diff, st])
    wb.save(out_xlsx)
    print(f"done: {len(a_map)} agents (A), {len(agg)} agents (B); summary written to {out_xlsx}")
    return 0

if __name__ == "__main__":
    sys.exit(main())
